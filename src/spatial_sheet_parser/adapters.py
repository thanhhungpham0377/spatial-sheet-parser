"""Input adapters for TSV, CSV, JSON grid, raw 2D lists, Excel (.xlsx/.xls), and OpenDocument (.ods) workbooks."""

import base64
import csv
import io
import json
import os
import zipfile
from typing import Any, Dict, List, Optional, Union
from spatial_sheet_parser.grid import Grid

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import xlrd
except ImportError:
    xlrd = None

try:
    import odf.opendocument
    import odf.table
    import odf.text
except ImportError:
    odf = None


class TSVAdapter:
    """Adapter for TSV text input."""

    @staticmethod
    def parse(
        content: str,
        sheet_name: Optional[str] = None,
        merged_cell_policy: str = "top_left_only",
    ) -> Grid:
        if content.startswith("\ufeff"):
            content = content[1:]
        matrix: List[List[Any]] = []
        try:
            reader = csv.reader(io.StringIO(content), delimiter="\t")
            for row in reader:
                matrix.append(row)
        except Exception:
            for line in content.splitlines():
                matrix.append(line.split("\t"))
        return Grid.from_matrix(
            matrix,
            source_format="tsv",
            sheet_name=sheet_name,
            merged_cell_policy=merged_cell_policy,
        )



class CSVAdapter:
    """Adapter for CSV text input."""

    @staticmethod
    def parse(
        content: str,
        sheet_name: Optional[str] = None,
        merged_cell_policy: str = "top_left_only",
    ) -> Grid:
        if content.startswith("\ufeff"):
            content = content[1:]
        matrix: List[List[Any]] = []
        reader = csv.reader(io.StringIO(content))
        for row in reader:
            matrix.append(row)
        return Grid.from_matrix(
            matrix,
            source_format="csv",
            sheet_name=sheet_name,
            merged_cell_policy=merged_cell_policy,
        )


class JSONGridAdapter:
    """Adapter for JSON grid input."""

    @staticmethod
    def parse(
        content: Union[str, Dict[str, Any], List[List[Any]]],
        sheet_name: Optional[str] = None,
        merged_cell_policy: str = "top_left_only",
    ) -> Grid:
        if isinstance(content, str):
            if content.startswith("\ufeff"):
                content = content[1:]
            parsed = json.loads(content)
        else:
            parsed = content

        merged_cells: Optional[List[Dict[str, int]]] = None
        extracted_sheet_name = sheet_name

        if isinstance(parsed, dict):
            matrix = parsed.get("grid", parsed.get("data", []))
            merged_cells = parsed.get("merged_cells", None)
            extracted_sheet_name = parsed.get("sheet_name", sheet_name)
        elif isinstance(parsed, list):
            matrix = parsed
        else:
            raise ValueError(f"Invalid JSON grid format: {type(parsed)}")

        return Grid.from_matrix(
            matrix,
            source_format="json",
            sheet_name=extracted_sheet_name,
            merged_cells=merged_cells,
            merged_cell_policy=merged_cell_policy,
        )


class ExcelAdapter:
    """Adapter for Excel (.xlsx) workbook input using openpyxl."""

    @staticmethod
    def parse(
        content: Union[str, bytes, io.BytesIO],
        sheet_name: Optional[str] = None,
        merged_cell_policy: str = "top_left_only",
    ) -> Grid:
        if openpyxl is None:
            raise ImportError(
                "openpyxl package is required for .xlsx support. "
                "Install it with `pip install openpyxl>=3.1.0`."
            )

        file_like: Any = None
        if isinstance(content, bytes):
            file_like = io.BytesIO(content)
        elif isinstance(content, io.BytesIO):
            file_like = content
        elif isinstance(content, str):
            if os.path.exists(content) and os.path.isfile(content):
                file_like = content
            else:
                clean_str = content.split(",", 1)[-1] if content.startswith("data:") else content
                try:
                    decoded = base64.b64decode(clean_str)
                    if decoded.startswith(b"PK\x03\x04"):
                        file_like = io.BytesIO(decoded)
                except Exception:
                    pass

        if file_like is None:
            raise ValueError("Invalid Excel workbook input or file not found.")

        try:
            wb = openpyxl.load_workbook(file_like, data_only=True)
        except Exception as e:
            raise ValueError(f"Failed to parse Excel workbook: {e}")

        available_sheets = wb.sheetnames
        if not available_sheets:
            raise ValueError("Excel workbook contains no worksheets.")

        if sheet_name:
            if sheet_name in available_sheets:
                ws = wb[sheet_name]
                selected_sheet = sheet_name
            else:
                raise ValueError(
                    f"Worksheet '{sheet_name}' not found in workbook. Available sheets: {available_sheets}"
                )
        else:
            ws = wb.active if wb.active is not None else wb.worksheets[0]
            selected_sheet = ws.title

        merged_cells: List[Dict[str, int]] = []
        for mrange in ws.merged_cells.ranges:
            merged_cells.append({
                "start_row": mrange.min_row - 1,
                "start_col": mrange.min_col - 1,
                "end_row": mrange.max_row - 1,
                "end_col": mrange.max_col - 1,
            })

        matrix: List[List[Any]] = []
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0

        for r in range(1, max_row + 1):
            row_vals: List[Any] = []
            for c in range(1, max_col + 1):
                val = ws.cell(row=r, column=c).value
                if hasattr(val, "isoformat"):
                    val = val.isoformat()
                row_vals.append(val)
            matrix.append(row_vals)

        return Grid.from_matrix(
            matrix,
            source_format="xlsx",
            sheet_name=selected_sheet,
            merged_cells=merged_cells,
            merged_cell_policy=merged_cell_policy,
        )

    @staticmethod
    def get_sheet_names(content: Union[str, bytes, io.BytesIO]) -> List[str]:
        if openpyxl is None:
            raise ImportError(
                "openpyxl package is required for .xlsx support. "
                "Install it with `pip install openpyxl>=3.1.0`."
            )

        file_like: Any = None
        if isinstance(content, bytes):
            file_like = io.BytesIO(content)
        elif isinstance(content, io.BytesIO):
            file_like = content
        elif isinstance(content, str):
            if os.path.exists(content) and os.path.isfile(content):
                file_like = content
            else:
                clean_str = content.split(",", 1)[-1] if content.startswith("data:") else content
                try:
                    decoded = base64.b64decode(clean_str)
                    if decoded.startswith(b"PK\x03\x04"):
                        file_like = io.BytesIO(decoded)
                except Exception:
                    pass

        if file_like is None:
            raise ValueError("Invalid Excel workbook input or file not found.")

        try:
            wb = openpyxl.load_workbook(file_like, read_only=True)
            names = list(wb.sheetnames)
            wb.close()
            return names
        except Exception as e:
            raise ValueError(f"Failed to inspect Excel workbook sheets: {e}")



class XLSAdapter:
    """Adapter for legacy Excel (.xls) workbook input using xlrd."""

    @staticmethod
    def parse(
        content: Union[str, bytes, io.BytesIO],
        sheet_name: Optional[str] = None,
        merged_cell_policy: str = "top_left_only",
    ) -> Grid:
        if xlrd is None:
            raise ImportError(
                "xlrd package is required for legacy Excel (.xls) support. "
                "Install it with `pip install xlrd>=2.0.1`."
            )

        file_bytes: Optional[bytes] = None
        if isinstance(content, bytes):
            file_bytes = content
        elif isinstance(content, io.BytesIO):
            file_bytes = content.getvalue()
        elif isinstance(content, str):
            if os.path.exists(content) and os.path.isfile(content):
                with open(content, "rb") as f:
                    file_bytes = f.read()
            else:
                clean_str = content.split(",", 1)[-1] if content.startswith("data:") else content
                try:
                    decoded = base64.b64decode(clean_str)
                    if decoded.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
                        file_bytes = decoded
                except Exception:
                    pass

        if file_bytes is None:
            raise ValueError("Invalid XLS workbook input or file not found.")

        try:
            try:
                wb = xlrd.open_workbook(file_contents=file_bytes, formatting_info=True)
                has_formatting = True
            except Exception:
                wb = xlrd.open_workbook(file_contents=file_bytes, formatting_info=False)
                has_formatting = False
        except Exception as e:
            raise ValueError(f"Failed to parse XLS workbook: {e}")

        available_sheets = wb.sheet_names()
        if not available_sheets:
            raise ValueError("XLS workbook contains no worksheets.")

        if sheet_name:
            if sheet_name in available_sheets:
                ws = wb.sheet_by_name(sheet_name)
                selected_sheet = sheet_name
            else:
                raise ValueError(
                    f"Worksheet '{sheet_name}' not found in workbook. Available sheets: {available_sheets}"
                )
        else:
            ws = wb.sheet_by_index(0)
            selected_sheet = ws.name

        merged_cells: List[Dict[str, int]] = []
        if has_formatting and hasattr(ws, "merged_cells"):
            for rlow, rhigh, clow, chigh in ws.merged_cells:
                merged_cells.append({
                    "start_row": rlow,
                    "start_col": clow,
                    "end_row": rhigh - 1,
                    "end_col": chigh - 1,
                })

        matrix: List[List[Any]] = []
        for r in range(ws.nrows):
            row_vals: List[Any] = []
            for c in range(ws.ncols):
                val = ws.cell_value(r, c)
                if val == "":
                    val = None
                row_vals.append(val)
            matrix.append(row_vals)

        return Grid.from_matrix(
            matrix,
            source_format="xls",
            sheet_name=selected_sheet,
            merged_cells=merged_cells,
            merged_cell_policy=merged_cell_policy,
        )

    @staticmethod
    def get_sheet_names(content: Union[str, bytes, io.BytesIO]) -> List[str]:
        if xlrd is None:
            raise ImportError(
                "xlrd package is required for legacy Excel (.xls) support. "
                "Install it with `pip install xlrd>=2.0.1`."
            )

        file_bytes: Optional[bytes] = None
        if isinstance(content, bytes):
            file_bytes = content
        elif isinstance(content, io.BytesIO):
            file_bytes = content.getvalue()
        elif isinstance(content, str):
            if os.path.exists(content) and os.path.isfile(content):
                with open(content, "rb") as f:
                    file_bytes = f.read()
            else:
                clean_str = content.split(",", 1)[-1] if content.startswith("data:") else content
                try:
                    decoded = base64.b64decode(clean_str)
                    if decoded.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
                        file_bytes = decoded
                except Exception:
                    pass

        if file_bytes is None:
            raise ValueError("Invalid XLS workbook input or file not found.")

        try:
            wb = xlrd.open_workbook(file_contents=file_bytes, on_demand=True)
            return list(wb.sheet_names())
        except Exception as e:
            raise ValueError(f"Failed to inspect XLS workbook sheets: {e}")



class ODSAdapter:
    """Adapter for OpenDocument Spreadsheet (.ods) workbook input using odfpy."""

    @staticmethod
    def _get_text(elem: Any) -> str:
        texts: List[str] = []
        for child in getattr(elem, "childNodes", []):
            if getattr(child, "nodeType", None) == getattr(child, "TEXT_NODE", 3):
                texts.append(child.data)
            elif hasattr(child, "childNodes"):
                texts.append(ODSAdapter._get_text(child))
        return "".join(texts)

    @staticmethod
    def parse(
        content: Union[str, bytes, io.BytesIO],
        sheet_name: Optional[str] = None,
        merged_cell_policy: str = "top_left_only",
    ) -> Grid:
        if odf is None:
            raise ImportError(
                "odfpy package is required for OpenDocument Spreadsheet (.ods) support. "
                "Install it with `pip install odfpy>=1.4.1`."
            )

        file_like: Any = None
        if isinstance(content, bytes):
            file_like = io.BytesIO(content)
        elif isinstance(content, io.BytesIO):
            file_like = content
        elif isinstance(content, str):
            if os.path.exists(content) and os.path.isfile(content):
                file_like = content
            else:
                clean_str = content.split(",", 1)[-1] if content.startswith("data:") else content
                try:
                    decoded = base64.b64decode(clean_str)
                    if decoded.startswith(b"PK\x03\x04"):
                        file_like = io.BytesIO(decoded)
                except Exception:
                    pass

        if file_like is None:
            raise ValueError("Invalid ODS workbook input or file not found.")

        try:
            doc = odf.opendocument.load(file_like)
        except Exception as e:
            raise ValueError(f"Failed to parse ODS workbook: {e}")

        tables = doc.getElementsByType(odf.table.Table)
        if not tables:
            raise ValueError("ODS workbook contains no worksheets.")

        available_sheets = [t.getAttribute("name") or f"Sheet{i+1}" for i, t in enumerate(tables)]

        if sheet_name:
            if sheet_name in available_sheets:
                target_table = tables[available_sheets.index(sheet_name)]
                selected_sheet = sheet_name
            else:
                raise ValueError(
                    f"Worksheet '{sheet_name}' not found in workbook. Available sheets: {available_sheets}"
                )
        else:
            target_table = tables[0]
            selected_sheet = available_sheets[0]

        raw_rows = target_table.getElementsByType(odf.table.TableRow)
        matrix: List[List[Any]] = []
        merged_cells: List[Dict[str, int]] = []

        current_r = 0
        for row_elem in raw_rows:
            row_rep = 1
            if hasattr(row_elem, "attributes"):
                for (ns, name), val in row_elem.attributes.items():
                    if name == "number-rows-repeated":
                        try:
                            row_rep = int(val)
                        except ValueError:
                            pass

            row_cells: List[Any] = []
            c_idx = 0
            for child in getattr(row_elem, "childNodes", []):
                if not hasattr(child, "attributes"):
                    continue
                tag = getattr(child, "tagName", "")
                if "table-cell" not in tag and "covered-table-cell" not in tag:
                    continue

                col_span = 1
                row_span = 1
                col_rep = 1
                for (ns, attr_name), val in child.attributes.items():
                    if attr_name == "number-columns-spanned":
                        try:
                            col_span = int(val)
                        except ValueError:
                            pass
                    elif attr_name == "number-rows-spanned":
                        try:
                            row_span = int(val)
                        except ValueError:
                            pass
                    elif attr_name == "number-columns-repeated":
                        try:
                            col_rep = int(val)
                        except ValueError:
                            pass

                val_text = ODSAdapter._get_text(child).strip()
                if not val_text:
                    val_text = None

                if (col_span > 1 or row_span > 1) and "covered-table-cell" not in tag:
                    merged_cells.append({
                        "start_row": current_r,
                        "start_col": c_idx,
                        "end_row": current_r + row_span - 1,
                        "end_col": c_idx + col_span - 1,
                    })

                if val_text is None and col_rep > 100:
                    col_rep = 1

                for _ in range(col_rep):
                    row_cells.append(val_text)
                    c_idx += 1

            if not any(v is not None for v in row_cells) and row_rep > 100:
                row_rep = 1

            for _ in range(row_rep):
                matrix.append(list(row_cells))
                current_r += 1

        while matrix and not any(v is not None for v in matrix[-1]):
            matrix.pop()

        if matrix:
            last_non_empty_col = -1
            for r in matrix:
                for c in range(len(r) - 1, last_non_empty_col, -1):
                    if r[c] is not None:
                        last_non_empty_col = c

            trim_col = max(0, last_non_empty_col + 1)
            matrix = [r[:trim_col] for r in matrix]

        return Grid.from_matrix(
            matrix,
            source_format="ods",
            sheet_name=selected_sheet,
            merged_cells=merged_cells,
            merged_cell_policy=merged_cell_policy,
        )

    @staticmethod
    def get_sheet_names(content: Union[str, bytes, io.BytesIO]) -> List[str]:
        if odf is None:
            raise ImportError(
                "odfpy package is required for OpenDocument Spreadsheet (.ods) support. "
                "Install it with `pip install odfpy>=1.4.1`."
            )

        file_like: Any = None
        if isinstance(content, bytes):
            file_like = io.BytesIO(content)
        elif isinstance(content, io.BytesIO):
            file_like = content
        elif isinstance(content, str):
            if os.path.exists(content) and os.path.isfile(content):
                file_like = content
            else:
                clean_str = content.split(",", 1)[-1] if content.startswith("data:") else content
                try:
                    decoded = base64.b64decode(clean_str)
                    if decoded.startswith(b"PK\x03\x04"):
                        file_like = io.BytesIO(decoded)
                except Exception:
                    pass

        if file_like is None:
            raise ValueError("Invalid ODS workbook input or file not found.")

        try:
            doc = odf.opendocument.load(file_like)
            tables = doc.getElementsByType(odf.table.Table)
            return [t.getAttribute("name") or f"Sheet{i+1}" for i, t in enumerate(tables)]
        except Exception as e:
            raise ValueError(f"Failed to inspect ODS workbook sheets: {e}")



def _detect_zip_format(raw_bytes: bytes) -> str:
    """Inspects zip container files to distinguish between XLSX and ODS formats."""
    try:
        with zipfile.ZipFile(io.BytesIO(raw_bytes)) as zf:
            namelist = zf.namelist()
            if "content.xml" in namelist or "META-INF/manifest.xml" in namelist:
                return "ods"
            if any(f.startswith("xl/") for f in namelist) or "[Content_Types].xml" in namelist:
                return "xlsx"
    except Exception:
        pass
    return "xlsx"


def inspect_workbook_sheets(
    content: Union[str, bytes, io.BytesIO],
    source_format: str = "auto",
) -> List[str]:
    """Inspects workbook content and returns list of sheet names."""
    fmt = source_format.lower()
    if fmt in ("xlsx", "excel", "xlsm", "xltx", "xltm"):
        return ExcelAdapter.get_sheet_names(content)
    if fmt == "xls":
        return XLSAdapter.get_sheet_names(content)
    if fmt == "ods":
        return ODSAdapter.get_sheet_names(content)

    file_bytes: Optional[bytes] = None
    if isinstance(content, bytes):
        file_bytes = content
    elif isinstance(content, io.BytesIO):
        file_bytes = content.getvalue()
    elif isinstance(content, str):
        if os.path.exists(content) and os.path.isfile(content):
            ext = os.path.splitext(content)[1].lower()
            if ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
                return ExcelAdapter.get_sheet_names(content)
            elif ext == ".xls":
                return XLSAdapter.get_sheet_names(content)
            elif ext == ".ods":
                return ODSAdapter.get_sheet_names(content)
            return []
        else:
            clean_str = content.split(",", 1)[-1] if content.startswith("data:") else content
            try:
                decoded = base64.b64decode(clean_str)
                file_bytes = decoded
            except Exception:
                pass

    if file_bytes:
        if file_bytes.startswith(b"PK\x03\x04"):
            detected = _detect_zip_format(file_bytes)
            if detected == "ods":
                return ODSAdapter.get_sheet_names(file_bytes)
            else:
                return ExcelAdapter.get_sheet_names(file_bytes)
        elif file_bytes.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            return XLSAdapter.get_sheet_names(file_bytes)

    return []


def normalize_input(
    grid_input: Union[List[List[Any]], str, bytes, Dict[str, Any], Grid],
    source_format: str = "raw",
    sheet_name: Optional[str] = None,
    merged_cell_policy: str = "top_left_only",
) -> Grid:
    """Normalizes arbitrary grid input into a standard Grid instance."""
    if isinstance(grid_input, Grid):
        return grid_input

    fmt = source_format.lower()

    if fmt in ("xlsx", "excel", "xlsm", "xltx", "xltm"):
        return ExcelAdapter.parse(grid_input, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)

    if fmt == "xls":
        return XLSAdapter.parse(grid_input, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)

    if fmt == "ods":
        return ODSAdapter.parse(grid_input, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)

    if fmt == "tsv":
        if isinstance(grid_input, bytes):
            grid_input = grid_input.decode("utf-8-sig")
        if isinstance(grid_input, str):
            return TSVAdapter.parse(grid_input, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)

    if fmt == "csv":
        if isinstance(grid_input, bytes):
            grid_input = grid_input.decode("utf-8-sig")
        if isinstance(grid_input, str):
            return CSVAdapter.parse(grid_input, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)

    if fmt == "json" or isinstance(grid_input, dict):
        if isinstance(grid_input, bytes):
            grid_input = grid_input.decode("utf-8-sig")
        return JSONGridAdapter.parse(grid_input, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)

    # Auto-detection for raw/auto format or arbitrary inputs
    if isinstance(grid_input, bytes):
        if grid_input.startswith(b"PK\x03\x04"):
            detected = _detect_zip_format(grid_input)
            if detected == "ods":
                return ODSAdapter.parse(grid_input, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)
            else:
                return ExcelAdapter.parse(grid_input, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)
        elif grid_input.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            return XLSAdapter.parse(grid_input, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)
        else:
            try:
                grid_input = grid_input.decode("utf-8-sig")
            except UnicodeDecodeError:
                raise ValueError("Unable to decode binary spreadsheet input. Unrecognized binary format.")

    if isinstance(grid_input, str):
        low_str = grid_input.lower()
        if os.path.exists(grid_input) and os.path.isfile(grid_input):
            ext = os.path.splitext(grid_input)[1].lower()
            if ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
                return ExcelAdapter.parse(grid_input, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)
            elif ext == ".xls":
                return XLSAdapter.parse(grid_input, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)
            elif ext == ".ods":
                return ODSAdapter.parse(grid_input, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)
            elif ext == ".tsv":
                with open(grid_input, "r", encoding="utf-8-sig") as f:
                    return TSVAdapter.parse(f.read(), sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)
            elif ext == ".csv":
                with open(grid_input, "r", encoding="utf-8-sig") as f:
                    return CSVAdapter.parse(f.read(), sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)
            elif ext == ".json":
                with open(grid_input, "r", encoding="utf-8-sig") as f:
                    return JSONGridAdapter.parse(f.read(), sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)

        clean_str = grid_input.split(",", 1)[-1] if grid_input.startswith("data:") else grid_input
        try:
            decoded = base64.b64decode(clean_str)
            if decoded.startswith(b"PK\x03\x04"):
                detected = _detect_zip_format(decoded)
                if detected == "ods":
                    return ODSAdapter.parse(decoded, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)
                else:
                    return ExcelAdapter.parse(decoded, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)
            elif decoded.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
                return XLSAdapter.parse(decoded, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)
        except Exception:
            pass

        stripped = grid_input.strip()
        if stripped.startswith("{") or stripped.startswith("["):
            try:
                return JSONGridAdapter.parse(grid_input, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)
            except Exception:
                pass

        if "\t" in grid_input:
            return TSVAdapter.parse(grid_input, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)

        if "," in grid_input:
            return CSVAdapter.parse(grid_input, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)

        return TSVAdapter.parse(grid_input, sheet_name=sheet_name, merged_cell_policy=merged_cell_policy)

    if isinstance(grid_input, list):
        return Grid.from_matrix(
            grid_input,
            source_format=source_format if source_format != "auto" else "raw",
            sheet_name=sheet_name,
            merged_cell_policy=merged_cell_policy,
        )

    raise ValueError(f"Unsupported grid input type: {type(grid_input)}")
