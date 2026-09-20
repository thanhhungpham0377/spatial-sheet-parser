"""Unit tests for Excel (.xlsx) adapter, sheet selection, merged cells, and error handling."""

import io
import os
import sys
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "src")))

import openpyxl
from spatial_sheet_parser import SpatialSheetParser, ParserConfig, Grid
from spatial_sheet_parser.adapters import ExcelAdapter, normalize_input, inspect_workbook_sheets


def create_test_workbook(xlsx_path: str):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sales Overview"
    ws1["A1"] = "Sales Overview 2026"
    ws1.merge_cells("A1:C1")
    ws1["B2"] = "Product"
    ws1["C2"] = "Revenue"
    ws1["B3"] = "Widget Alpha"
    ws1["C3"] = 1500
    ws1["B4"] = "Widget Beta"
    ws1["C4"] = 2500

    ws2 = wb.create_sheet("Inventory")
    ws2["A1"] = "Warehouse Stock"
    ws2["B2"] = "SKU"
    ws2["C2"] = "Quantity"
    ws2["B3"] = "SKU-001"
    ws2["C3"] = 500
    wb.save(xlsx_path)


class TestExcelAdapter(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.fixture_dir = os.path.join(os.path.dirname(__file__), "fixtures")
        cls.xlsx_path = os.path.join(cls.fixture_dir, "sample_workbook.xlsx")
        create_test_workbook(cls.xlsx_path)

    def test_excel_adapter_default_sheet(self):
        grid = ExcelAdapter.parse(self.xlsx_path)
        self.assertIsInstance(grid, Grid)
        self.assertEqual(grid.source_format, "xlsx")
        self.assertEqual(grid.sheet_name, "Sales Overview")
        self.assertEqual(grid.get_value(0, 0), "Sales Overview 2026")
        self.assertEqual(grid.get_value(1, 1), "Product")
        self.assertEqual(grid.get_value(2, 2), 1500)

    def test_excel_adapter_explicit_sheet_selection(self):
        grid = ExcelAdapter.parse(self.xlsx_path, sheet_name="Inventory")
        self.assertEqual(grid.sheet_name, "Inventory")
        self.assertEqual(grid.get_value(0, 0), "Warehouse Stock")
        self.assertEqual(grid.get_value(1, 1), "SKU")
        self.assertEqual(grid.get_value(2, 2), 500)

    def test_excel_adapter_missing_sheet_error(self):
        with self.assertRaises(ValueError) as ctx:
            ExcelAdapter.parse(self.xlsx_path, sheet_name="NonExistentSheet")
        err_msg = str(ctx.exception)
        self.assertIn("Worksheet 'NonExistentSheet' not found", err_msg)
        self.assertIn("Sales Overview", err_msg)
        self.assertIn("Inventory", err_msg)

    def test_excel_adapter_merged_cells_top_left_policy(self):
        grid = ExcelAdapter.parse(self.xlsx_path)
        top_left = grid.get_cell(0, 0)
        self.assertIsNotNone(top_left)
        self.assertEqual(top_left.value, "Sales Overview 2026")
        self.assertTrue(top_left.is_top_left)
        self.assertTrue(top_left.merged)

        cell_b1 = grid.get_cell(0, 1)
        self.assertIsNotNone(cell_b1)
        self.assertIsNone(cell_b1.value)
        self.assertFalse(cell_b1.is_top_left)
        self.assertTrue(cell_b1.merged)

    def test_excel_adapter_binary_bytes_input(self):
        with open(self.xlsx_path, "rb") as f:
            bytes_data = f.read()
        grid = ExcelAdapter.parse(bytes_data, sheet_name="Sales Overview")
        self.assertEqual(grid.sheet_name, "Sales Overview")
        self.assertEqual(grid.get_value(0, 0), "Sales Overview 2026")

    def test_excel_adapter_malformed_workbook(self):
        bad_bytes = b"PK\x03\x04this is corrupted binary ZIP content"
        with self.assertRaises(ValueError) as ctx:
            ExcelAdapter.parse(bad_bytes)
        self.assertIn("Failed to parse Excel workbook", str(ctx.exception))

    def test_xls_support_and_parsing(self):
        xls_fixture = os.path.join(self.fixture_dir, "sample_workbook.xls")
        grid = normalize_input(xls_fixture, source_format="xls")
        self.assertEqual(grid.source_format, "xls")
        self.assertEqual(grid.sheet_name, "SalesSummary")

        with self.assertRaises(ValueError) as ctx:
            normalize_input(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1BAD_XLS_BYTES", source_format="xls")
        self.assertIn("Failed to parse XLS workbook", str(ctx.exception))

    def test_parser_core_integration_with_excel(self):
        parser = SpatialSheetParser()
        result = parser.parse_grid(self.xlsx_path, source_format="xlsx", sheet_name="Inventory")
        self.assertEqual(result["metadata"]["source_format"], "xlsx")
        self.assertEqual(result["metadata"]["sheet_name"], "Inventory")
        self.assertTrue(result["metadata"]["validation"]["is_valid"])
        self.assertGreaterEqual(len(result["tables"]), 1)
        self.assertEqual(result["tables"][0]["title"], "Warehouse Stock")

    def test_excel_get_sheet_names(self):
        sheets = ExcelAdapter.get_sheet_names(self.xlsx_path)
        self.assertEqual(sheets, ["Sales Overview", "Inventory"])

        with open(self.xlsx_path, "rb") as f:
            bytes_data = f.read()
        sheets_bytes = ExcelAdapter.get_sheet_names(bytes_data)
        self.assertEqual(sheets_bytes, ["Sales Overview", "Inventory"])

    def test_inspect_workbook_sheets(self):
        sheets = inspect_workbook_sheets(self.xlsx_path, source_format="xlsx")
        self.assertEqual(sheets, ["Sales Overview", "Inventory"])

        xls_fixture = os.path.join(self.fixture_dir, "sample_workbook.xls")
        xls_sheets = inspect_workbook_sheets(xls_fixture, source_format="xls")
        self.assertEqual(xls_sheets, ["SalesSummary", "Inventory"])



if __name__ == "__main__":

    unittest.main()
